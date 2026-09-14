import os
import re
import sys
import json
import sqlite3
import shutil
import tempfile
import threading
import subprocess
import urllib.error
import urllib.request
from pathlib import Path


def app_dir():
    """Папка, где лежит exe или исходный .py."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def data_dir():
    """Служебные файлы обновления — не рядом с exe."""
    if os.name == "nt":
        base = os.environ.get("LOCALAPPDATA") or os.path.join(
            os.path.expanduser("~"), "AppData", "Local"
        )
    else:
        base = os.environ.get("XDG_DATA_HOME") or os.path.join(
            os.path.expanduser("~"), ".local", "share"
        )
    path = os.path.join(base, "ProductManager")
    os.makedirs(path, exist_ok=True)
    return path


def runtime_dir():
    path = os.path.join(data_dir(), "runtime")
    os.makedirs(path, exist_ok=True)
    return path


def _find_lib_dir(root, required, max_depth=4):
    if not root or not os.path.isdir(root):
        return None
    root = os.path.abspath(root)
    for dirpath, dirnames, filenames in os.walk(root):
        if all(name in filenames for name in required):
            return dirpath
        rel = os.path.relpath(dirpath, root)
        depth = 0 if rel == "." else rel.count(os.sep) + 1
        if depth >= max_depth:
            dirnames[:] = []
    return None


def persist_runtime():
    """Копирует Tcl/Tk из текущей распаковки exe в AppData."""
    if not getattr(sys, "frozen", False):
        return
    src_root = getattr(sys, "_MEIPASS", None)
    if not src_root or not os.path.isdir(src_root):
        return
    dest_root = runtime_dir()
    names = (
        "_tcl_data",
        "_tk_data",
        "tcl",
        "tk",
        "tcl8",
        "tcl8.6",
        "tk8.6",
    )
    for name in names:
        src = os.path.join(src_root, name)
        dest = os.path.join(dest_root, name)
        if not os.path.isdir(src):
            continue
        try:
            if os.path.isdir(dest) and _find_lib_dir(dest, ("init.tcl",), max_depth=3):
                continue
            if os.path.isdir(dest):
                shutil.rmtree(dest, ignore_errors=True)
            shutil.copytree(src, dest)
        except OSError:
            pass


import tkinter as tk
from tkinter import ttk, messagebox, filedialog

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
except ImportError:
    DND_FILES = None
    TkinterDnD = None


def cleanup_app_folder():
    """Удаляет служебные файлы обновления из папки с exe."""
    hidden = data_dir()
    folder = app_dir()

    for name in (LOCAL_UPDATE_SCRIPT, "version.txt", "update_repo.txt"):
        src = os.path.join(folder, name)
        dst = os.path.join(hidden, name)
        if os.path.isfile(src) and not os.path.isfile(dst):
            try:
                os.replace(src, dst)
            except OSError:
                try:
                    with open(src, "rb") as fsrc, open(dst, "wb") as fdst:
                        fdst.write(fsrc.read())
                except OSError:
                    pass

    db_src = os.path.join(folder, "products.db")
    db_dst = os.path.join(hidden, "products.db")
    if os.path.isfile(db_src) and not os.path.isfile(db_dst):
        try:
            shutil.copy2(db_src, db_dst)
        except OSError:
            pass

    leftovers = (
        "app_latest.py",
        "app_latest.py.tmp",
        "app_latest.py.bad",
        "version.txt",
        "apply_update.bat",
        "ProductManager_new.exe",
        "Tovary_new.exe",
        "update_repo.txt",
    )
    for name in leftovers:
        path = os.path.join(folder, name)
        try:
            if os.path.isfile(path):
                os.remove(path)
        except OSError:
            pass


def resource_path(name):
    """Файл из сборки PyInstaller (иконка и т.п.)."""
    if getattr(sys, "frozen", False):
        base = getattr(sys, "_MEIPASS", app_dir())
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, name)


def resolve_db_file():
    dest = os.path.join(data_dir(), "products.db")
    src = os.path.join(app_dir(), "products.db")
    if os.path.isfile(src) and os.path.abspath(src) != os.path.abspath(dest):
        if not os.path.isfile(dest):
            try:
                shutil.copy2(src, dest)
            except OSError:
                return src
    return dest


DB_FILE = resolve_db_file()

# Репозиторий GitHub в формате владелец/имя.
# Можно также положить рядом с exe файл update_repo.txt с этой строкой.
GITHUB_REPO = "Igorgawrilin/kaspi_sort_exel"
GITHUB_BRANCH = "gh_pages"
UPDATE_SCRIPT_NAME = "product_manager_sets.py"
LOCAL_UPDATE_SCRIPT = "app_latest.py"
APP_NAME = "Product Manager"


def _version_from_file(path):
    if not path or not os.path.isfile(path):
        return ""
    try:
        with open(path, encoding="utf-8") as f:
            return f.read().strip().splitlines()[0].strip()
    except OSError:
        return ""


def bundled_app_version():
    return _version_from_file(resource_path("version.txt")) or "1.0.0"


def update_app_version():
    return _version_from_file(os.path.join(data_dir(), "version.txt"))


def read_app_version():
    bundled = bundled_app_version()
    updated = update_app_version()
    if updated and version_newer(updated, bundled):
        return updated
    if bundled:
        return bundled
    return updated or "1.0.0"


def get_github_repo():
    override = os.path.join(data_dir(), "update_repo.txt")
    if not os.path.isfile(override):
        override = os.path.join(app_dir(), "update_repo.txt")
    if os.path.isfile(override):
        try:
            with open(override, encoding="utf-8") as f:
                value = f.read().strip()
            if "/" in value and "YOUR_GITHUB" not in value:
                return value
        except OSError:
            pass
    return GITHUB_REPO


def parse_version(text):
    text = str(text or "").strip().lstrip("vV")
    parts = []
    for chunk in text.split("."):
        digits = ""
        for ch in chunk:
            if ch.isdigit():
                digits += ch
            else:
                break
        if digits == "":
            break
        parts.append(int(digits))
    return tuple(parts or [0])


def version_newer(latest, current):
    return parse_version(latest) > parse_version(current)


APP_VERSION = read_app_version()

try:
    from openpyxl import load_workbook, Workbook
except ImportError:
    load_workbook = None
    Workbook = None

try:
    import fitz
except ImportError:
    fitz = None


def extract_receipt_products(pdf_path):
    """Названия товаров из чека Kaspi: «Название .... 4 шт.»."""
    if fitz is None:
        return []
    try:
        document = fitz.open(pdf_path)
        full_text = ""
        for page in document:
            full_text += page.get_text("text") + "\n"
        document.close()
    except Exception:
        return []

    products = []
    for line in full_text.splitlines():
        line = line.strip()
        if not line:
            continue
        match = re.match(
            r"^(.*?)\s*(?:\.+\s*)+(\d+)\s*шт\.?\s*$",
            line,
            re.IGNORECASE,
        )
        if not match:
            continue
        product = re.sub(r"\s*\.+\s*$", "", match.group(1).strip())
        if product:
            products.append((product, int(match.group(2))))
    return products


def _folder_if_valid(value):
    if not value:
        return None
    text = str(value).strip().strip('"')
    if not text:
        return None
    path = Path(text)
    try:
        path = path.expanduser()
    except Exception:
        return None
    if path.is_dir():
        return path
    if path.is_file():
        return path.parent
    return None


def get_folder_from_clipboard(root=None):
    """Папка из буфера: Проводник (Ctrl+C) или текстовый путь."""
    if os.name == "nt":
        try:
            import ctypes
            from ctypes import wintypes

            user32 = ctypes.windll.user32
            shell32 = ctypes.windll.shell32
            CF_HDROP = 15
            user32.OpenClipboard.argtypes = [wintypes.HWND]
            user32.OpenClipboard.restype = wintypes.BOOL
            user32.GetClipboardData.argtypes = [wintypes.UINT]
            user32.GetClipboardData.restype = wintypes.HANDLE
            user32.CloseClipboard.restype = wintypes.BOOL
            shell32.DragQueryFileW.argtypes = [
                wintypes.HANDLE, wintypes.UINT, wintypes.LPWSTR, wintypes.UINT
            ]
            shell32.DragQueryFileW.restype = wintypes.UINT

            if user32.OpenClipboard(None):
                try:
                    handle = user32.GetClipboardData(CF_HDROP)
                    if handle:
                        count = shell32.DragQueryFileW(handle, 0xFFFFFFFF, None, 0)
                        if count:
                            length = shell32.DragQueryFileW(handle, 0, None, 0)
                            buf = ctypes.create_unicode_buffer(length + 2)
                            shell32.DragQueryFileW(handle, 0, buf, length + 2)
                            folder = _folder_if_valid(buf.value)
                            if folder:
                                return folder
                finally:
                    user32.CloseClipboard()
        except Exception:
            pass

    if root is not None:
        for kind in ("STRING", "UTF8_STRING", "TEXT"):
            try:
                folder = _folder_if_valid(root.clipboard_get(type=kind))
            except tk.TclError:
                folder = None
            if folder:
                return folder
        try:
            folder = _folder_if_valid(root.clipboard_get())
            if folder:
                return folder
        except tk.TclError:
            pass
    return None


class ProductApp:
    def __init__(self, root):
        self.root = root
        self.root.title(f"{APP_NAME}  v{APP_VERSION}")
        self.root.geometry("1100x700")
        self.root.minsize(950, 600)
        self.apply_app_icon()

        self.conn = sqlite3.connect(DB_FILE)
        self.create_db()

        self.external_entries = []
        self.excel_files = []
        self.excel_result = []
        self.pdf_selected_folder = None
        self.pdf_result_file = None
        self.pdf_files = []
        self.pdf_single = []
        self.pdf_multiple = []
        self.pdf_unknown = []
        self.pdf_processing = False
        self.latest_update = None
        self._update_busy = False

        self.setup_style()
        self.setup_clipboard()
        self.build_ui()
        self.load_products()
        try:
            persist_runtime()
        except Exception:
            pass

        self.root.protocol("WM_DELETE_WINDOW", self.close)
        self.root.after(800, self.check_for_updates_async)

    def apply_app_icon(self):
        icon_ico = resource_path("app.ico")
        if os.path.exists(icon_ico):
            try:
                self.root.iconbitmap(icon_ico)
            except tk.TclError:
                pass

        icon_png = resource_path("app.png")
        if os.path.exists(icon_png):
            try:
                self._app_icon = tk.PhotoImage(file=icon_png)
                self.root.iconphoto(True, self._app_icon)
            except tk.TclError:
                pass

    def create_db(self):
        cur = self.conn.cursor()
        cur.execute("""
            CREATE TABLE IF NOT EXISTS products (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                internal_sku TEXT NOT NULL UNIQUE
            )
        """)
        cur.execute("""
            CREATE TABLE IF NOT EXISTS external_skus (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                product_id INTEGER NOT NULL,
                sku TEXT NOT NULL,
                FOREIGN KEY(product_id) REFERENCES products(id) ON DELETE CASCADE,
                UNIQUE(product_id, sku)
            )
        """)
        self.conn.commit()

    def setup_style(self):
        style = ttk.Style()
        try:
            style.theme_use("vista")
        except tk.TclError:
            pass

        style.configure("Title.TLabel", font=("Segoe UI", 20, "bold"))
        style.configure("Subtitle.TLabel", font=("Segoe UI", 10))
        style.configure("Version.TLabel", font=("Segoe UI", 9), foreground="#666666")
        style.configure("Update.TLabel", font=("Segoe UI", 10, "bold"), foreground="#0a7a2f")
        style.configure("Big.TButton", font=("Segoe UI", 11, "bold"), padding=8)
        style.configure("Treeview", rowheight=30, font=("Segoe UI", 10))
        style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))

    def setup_clipboard(self):
        """Вставка и копирование в поля: Ctrl+C/V/X/A, Shift+Insert, правое меню."""

        def widget_of(event):
            widget = event.widget
            if isinstance(widget, str):
                try:
                    widget = self.root.nametowidget(widget)
                except tk.TclError:
                    return None
            try:
                if widget.winfo_class() in ("TEntry", "Entry"):
                    return widget
            except tk.TclError:
                return None
            return None

        def get_clipboard():
            for kind in ("STRING", "UTF8_STRING", "TEXT"):
                try:
                    value = self.root.clipboard_get(type=kind)
                    if value:
                        return value
                except tk.TclError:
                    pass
            try:
                return self.root.clipboard_get()
            except tk.TclError:
                return ""

        def delete_selection(widget):
            try:
                widget.delete("sel.first", "sel.last")
            except tk.TclError:
                pass

        def copy_text(widget):
            try:
                text = widget.selection_get()
            except tk.TclError:
                return
            try:
                self.root.clipboard_clear()
                self.root.clipboard_append(text)
                self.root.update_idletasks()
            except tk.TclError:
                pass

        def paste_text(widget):
            text = get_clipboard()
            if not text:
                return
            text = str(text).replace("\r\n", "\n").replace("\r", "\n")
            if "\n" in text:
                text = text.split("\n", 1)[0]
            delete_selection(widget)
            try:
                widget.insert("insert", text)
            except tk.TclError:
                try:
                    widget.event_generate("<<Paste>>")
                except tk.TclError:
                    pass

        def cut_text(widget):
            copy_text(widget)
            delete_selection(widget)

        def select_all(widget):
            try:
                widget.selection_range(0, "end")
                widget.icursor("end")
            except tk.TclError:
                pass

        def on_ctrl_key(event):
            widget = widget_of(event)
            if widget is None:
                return
            keysym = (event.keysym or "").lower()
            keycode = getattr(event, "keycode", 0)
            # На Windows keycode не зависит от русской раскладки: V=86, C=67, X=88, A=65
            if keycode == 86 or keysym == "v":
                paste_text(widget)
                return "break"
            if keycode == 67 or keysym == "c":
                copy_text(widget)
                return "break"
            if keycode == 88 or keysym == "x":
                cut_text(widget)
                return "break"
            if keycode == 65 or keysym == "a":
                select_all(widget)
                return "break"

        def on_shift_insert(event):
            widget = widget_of(event)
            if widget is None:
                return
            paste_text(widget)
            return "break"

        def show_menu(event):
            widget = widget_of(event)
            if widget is None:
                return
            widget.focus_set()
            menu = tk.Menu(widget, tearoff=0)
            menu.add_command(label="Вырезать", command=lambda: cut_text(widget))
            menu.add_command(label="Копировать", command=lambda: copy_text(widget))
            menu.add_command(label="Вставить", command=lambda: paste_text(widget))
            menu.add_separator()
            menu.add_command(label="Выделить всё", command=lambda: select_all(widget))
            try:
                menu.tk_popup(event.x_root, event.y_root)
            finally:
                menu.grab_release()
            return "break"

        for cls_name in ("TEntry", "Entry"):
            for sequence, handler in (
                ("<Control-KeyPress>", on_ctrl_key),
                ("<Shift-Insert>", on_shift_insert),
                ("<Button-3>", show_menu),
            ):
                try:
                    self.root.bind_class(cls_name, sequence, handler)
                except tk.TclError:
                    pass

    def build_ui(self):
        header = ttk.Frame(self.root, padding=(20, 18, 20, 10))
        header.pack(fill="x")

        title_row = ttk.Frame(header)
        title_row.pack(fill="x")

        ttk.Label(
            title_row, text=APP_NAME, style="Title.TLabel"
        ).pack(side="left", anchor="w")

        update_box = ttk.Frame(title_row)
        update_box.pack(side="right", anchor="e")

        self.version_label = ttk.Label(
            update_box,
            text=f"Версия {APP_VERSION}",
            style="Version.TLabel"
        )
        self.version_label.pack(side="left", padx=(0, 10))

        self.update_status_var = tk.StringVar(value="")
        self.update_status_label = ttk.Label(
            update_box,
            textvariable=self.update_status_var,
            style="Version.TLabel"
        )
        self.update_status_label.pack(side="left", padx=(0, 8))

        self.update_button = ttk.Button(
            update_box,
            text="Обновить",
            command=self.start_update,
            state="disabled"
        )
        self.update_button.pack(side="left", padx=(0, 6))

        ttk.Button(
            update_box,
            text="Проверить",
            command=self.check_for_updates_async
        ).pack(side="left")

        ttk.Label(
            header,
            text="Внутренний артикул + любое количество внешних артикулов",
            style="Subtitle.TLabel"
        ).pack(anchor="w", pady=(3, 0))

        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill="both", expand=True, padx=20, pady=(5, 20))

        products_tab = ttk.Frame(self.notebook, padding=15)
        self.notebook.add(products_tab, text="  Товары  ")
        self.build_products_tab(products_tab)

        excel_tab = ttk.Frame(self.notebook, padding=15)
        self.notebook.add(excel_tab, text="  Обработка Excel  ")
        self.build_excel_tab(excel_tab)

        self.pdf_tab = ttk.Frame(self.notebook, padding=15)
        self.notebook.add(self.pdf_tab, text="  Сортировка PDF  ")
        self.build_pdf_tab(self.pdf_tab)
        self.root.bind_all("<Control-v>", self.pdf_on_ctrl_v, add="+")
        self.root.bind_all("<Control-V>", self.pdf_on_ctrl_v, add="+")

    # -------------------- РАЗДЕЛ 1 --------------------

    def build_products_tab(self, parent):
        parent.columnconfigure(0, weight=1)
        parent.columnconfigure(1, weight=1)
        parent.rowconfigure(0, weight=1)

        form = ttk.LabelFrame(parent, text=" Создать товар ", padding=15)
        form.grid(row=0, column=0, sticky="nsew", padx=(0, 10))

        form.columnconfigure(1, weight=1)

        ttk.Label(form, text="Название товара:").grid(
            row=0, column=0, sticky="w", pady=7, padx=(0, 10)
        )
        self.name_var = tk.StringVar()
        self.name_entry = ttk.Entry(form, textvariable=self.name_var)
        self.name_entry.grid(row=0, column=1, sticky="ew", pady=7)

        ttk.Label(form, text="Внутренний артикул:").grid(
            row=1, column=0, sticky="w", pady=7, padx=(0, 10)
        )
        self.internal_var = tk.StringVar()
        ttk.Entry(form, textvariable=self.internal_var).grid(
            row=1, column=1, sticky="ew", pady=7
        )

        ttk.Label(form, text="Внешние артикулы:").grid(
            row=2, column=0, sticky="nw", pady=7, padx=(0, 10)
        )

        external_area = ttk.Frame(form)
        external_area.grid(row=2, column=1, sticky="nsew", pady=7)
        form.rowconfigure(2, weight=1)
        external_area.columnconfigure(0, weight=1)

        self.external_list = ttk.Frame(external_area)
        self.external_list.grid(row=0, column=0, sticky="nsew")

        ttk.Button(
            external_area,
            text="+ Добавить внешний артикул",
            command=self.add_external_entry
        ).grid(row=1, column=0, sticky="w", pady=(8, 0))

        self.add_external_entry()

        buttons = ttk.Frame(form)
        buttons.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(15, 0))

        ttk.Button(
            buttons, text="СОХРАНИТЬ", style="Big.TButton",
            command=self.save_product
        ).pack(side="left")

        ttk.Button(
            buttons, text="ОЧИСТИТЬ", command=self.clear_form
        ).pack(side="left", padx=(10, 0))

        table_frame = ttk.LabelFrame(parent, text=" Сохранённые товары ", padding=10)
        table_frame.grid(row=0, column=1, sticky="nsew", padx=(10, 0))
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(1, weight=1)

        search_frame = ttk.Frame(table_frame)
        search_frame.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        search_frame.columnconfigure(1, weight=1)

        ttk.Label(search_frame, text="Поиск:").grid(row=0, column=0, padx=(0, 8))
        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", lambda *_: self.load_products())
        ttk.Entry(search_frame, textvariable=self.search_var).grid(
            row=0, column=1, sticky="ew"
        )

        tree_frame = ttk.Frame(table_frame)
        tree_frame.grid(row=1, column=0, sticky="nsew")
        tree_frame.columnconfigure(0, weight=1)
        tree_frame.rowconfigure(0, weight=1)

        columns = ("id", "name", "internal", "external")
        self.tree = ttk.Treeview(tree_frame, columns=columns, show="headings")
        self.tree.heading("id", text="ID")
        self.tree.heading("name", text="Название")
        self.tree.heading("internal", text="Внутренний")
        self.tree.heading("external", text="Внешние артикулы")

        self.tree.column("id", width=45, anchor="center", stretch=False)
        self.tree.column("name", width=250)
        self.tree.column("internal", width=120)
        self.tree.column("external", width=220)

        scrollbar = ttk.Scrollbar(
            tree_frame, orient="vertical", command=self.tree.yview
        )
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.tree.bind("<Double-1>", self.show_product)

        actions = ttk.Frame(table_frame)
        actions.grid(row=2, column=0, sticky="ew", pady=(10, 0))

        ttk.Button(
            actions, text="Открыть / изменить", command=self.show_product
        ).pack(side="left")
        ttk.Button(
            actions, text="Удалить", command=self.delete_product
        ).pack(side="left", padx=(8, 0))

        ttk.Button(
            actions, text="Экспорт базы", command=self.export_products
        ).pack(side="right", padx=(8, 0))

        ttk.Button(
            actions, text="Импорт базы", command=self.import_products
        ).pack(side="right")

        self.status_var = tk.StringVar(value="Готово")
        ttk.Label(parent, textvariable=self.status_var).grid(
            row=1, column=0, columnspan=2, sticky="w", pady=(8, 0)
        )

    def export_products(self):
        """Экспортирует все товары и все их внешние артикулы в JSON-файл."""
        path = filedialog.asksaveasfilename(
            title="Сохранить базу товаров",
            defaultextension=".json",
            initialfile="База_товаров.json",
            filetypes=[("Файл базы товаров", "*.json"), ("Все файлы", "*.*")]
        )
        if not path:
            return

        cur = self.conn.cursor()
        cur.execute("SELECT id, name, internal_sku FROM products ORDER BY id")
        products = []

        for product_id, name, internal_sku in cur.fetchall():
            cur.execute(
                "SELECT sku FROM external_skus WHERE product_id = ? ORDER BY id",
                (product_id,)
            )
            external = [row[0] for row in cur.fetchall()]
            products.append({
                "name": name,
                "internal_sku": internal_sku,
                "external_skus": external
            })

        data = {
            "format": "ProductManagerDatabase",
            "version": 1,
            "products": products
        }

        try:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except OSError as exc:
            messagebox.showerror("Ошибка", f"Не удалось сохранить базу:\n\n{exc}")
            return

        self.status_var.set(f"База экспортирована: {len(products)} товаров")
        messagebox.showinfo(
            "База сохранена",
            f"Сохранено товаров: {len(products)}\n\nФайл:\n{path}\n\nЕго можно перенести на другое устройство и импортировать в программу."
        )

    def import_products(self):
        """Импортирует товары из JSON и объединяет их с текущей базой."""
        path = filedialog.askopenfilename(
            title="Загрузить базу товаров",
            filetypes=[("Файл базы товаров", "*.json"), ("Все файлы", "*.*")]
        )
        if not path:
            return

        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
        except (OSError, json.JSONDecodeError) as exc:
            messagebox.showerror("Ошибка", f"Не удалось открыть файл базы:\n\n{exc}")
            return

        if not isinstance(data, dict) or data.get("format") != "ProductManagerDatabase":
            messagebox.showerror(
                "Неверный файл",
                "Выбранный файл не является базой товаров этой программы."
            )
            return

        products = data.get("products")
        if not isinstance(products, list):
            messagebox.showerror("Ошибка", "В файле отсутствует список товаров.")
            return

        valid_products = []
        for item in products:
            if not isinstance(item, dict):
                continue
            name = str(item.get("name", "")).strip()
            internal = str(item.get("internal_sku", "")).strip()
            external = item.get("external_skus", [])
            if not name or not internal or not isinstance(external, list):
                continue
            external_clean = []
            for sku in external:
                sku = str(sku).strip()
                if sku and sku not in external_clean:
                    external_clean.append(sku)
            valid_products.append((name, internal, external_clean))

        if not valid_products:
            messagebox.showwarning("Пустая база", "В выбранном файле нет корректных товаров.")
            return

        confirm = messagebox.askyesno(
            "Импорт базы",
            f"Найдено товаров: {len(valid_products)}.\n\n"
            "Товары с уже существующим внутренним артикулом будут обновлены, "
            "а внешние артикулы объединены.\n\nПродолжить?"
        )
        if not confirm:
            return

        added = 0
        updated = 0
        external_added = 0
        errors = []

        try:
            cur = self.conn.cursor()
            for name, internal, external_skus in valid_products:
                cur.execute(
                    "SELECT id FROM products WHERE internal_sku = ?",
                    (internal,)
                )
                existing = cur.fetchone()

                if existing:
                    product_id = existing[0]
                    cur.execute(
                        "UPDATE products SET name = ? WHERE id = ?",
                        (name, product_id)
                    )
                    updated += 1
                else:
                    cur.execute(
                        "INSERT INTO products (name, internal_sku) VALUES (?, ?)",
                        (name, internal)
                    )
                    product_id = cur.lastrowid
                    added += 1

                for sku in external_skus:
                    try:
                        cur.execute(
                            "INSERT INTO external_skus (product_id, sku) VALUES (?, ?)",
                            (product_id, sku)
                        )
                        external_added += 1
                    except sqlite3.IntegrityError:
                        # Такой внешний артикул уже привязан к этому товару.
                        pass

            self.conn.commit()
        except sqlite3.Error as exc:
            self.conn.rollback()
            messagebox.showerror("Ошибка базы данных", f"Импорт отменён:\n\n{exc}")
            return

        self.load_products()
        self.status_var.set(
            f"Импорт завершён: добавлено {added}, обновлено {updated}, "
            f"добавлено внешних артикулов {external_added}"
        )
        messagebox.showinfo(
            "Импорт завершён",
            f"Добавлено новых товаров: {added}\n"
            f"Обновлено товаров: {updated}\n"
            f"Добавлено внешних артикулов: {external_added}"
        )

    def add_external_entry(self, value=""):
        row = ttk.Frame(self.external_list)
        row.pack(fill="x", pady=3)

        var = tk.StringVar(value=value)
        ttk.Entry(row, textvariable=var).pack(side="left", fill="x", expand=True)

        ttk.Button(
            row, text="×", width=3,
            command=lambda: self.remove_external_entry(row, var)
        ).pack(side="left", padx=(6, 0))

        self.external_entries.append((row, var))

    def remove_external_entry(self, row, var):
        if len(self.external_entries) <= 1:
            var.set("")
            return
        row.destroy()
        self.external_entries = [
            item for item in self.external_entries if item[0] != row
        ]

    def get_external_skus(self):
        result = []
        for _, var in self.external_entries:
            value = var.get().strip()
            if value and value not in result:
                result.append(value)
        return result

    def save_product(self):
        name = self.name_var.get().strip()
        internal = self.internal_var.get().strip()
        external = self.get_external_skus()

        if not name:
            messagebox.showwarning("Проверка", "Введите название товара.")
            self.name_entry.focus()
            return
        if not internal:
            messagebox.showwarning("Проверка", "Введите внутренний артикул.")
            return

        try:
            cur = self.conn.cursor()
            cur.execute(
                "INSERT INTO products (name, internal_sku) VALUES (?, ?)",
                (name, internal)
            )
            product_id = cur.lastrowid

            for sku in external:
                cur.execute(
                    "INSERT INTO external_skus (product_id, sku) VALUES (?, ?)",
                    (product_id, sku)
                )

            self.conn.commit()
            self.clear_form()
            self.load_products()
            self.status_var.set(f"Товар сохранён. ID: {product_id}")

        except sqlite3.IntegrityError:
            self.conn.rollback()
            messagebox.showerror(
                "Ошибка",
                f"Внутренний артикул «{internal}» уже существует."
            )

    def clear_form(self):
        self.name_var.set("")
        self.internal_var.set("")
        for row, _ in self.external_entries:
            row.destroy()
        self.external_entries = []
        self.add_external_entry()
        self.name_entry.focus()

    def load_products(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        search = self.search_var.get().strip() if hasattr(self, "search_var") else ""
        cur = self.conn.cursor()

        if search:
            pattern = f"%{search}%"
            cur.execute("""
                SELECT p.id, p.name, p.internal_sku,
                       GROUP_CONCAT(e.sku, ', ')
                FROM products p
                LEFT JOIN external_skus e ON e.product_id = p.id
                WHERE p.name LIKE ?
                   OR p.internal_sku LIKE ?
                   OR e.sku LIKE ?
                GROUP BY p.id
                ORDER BY p.id DESC
            """, (pattern, pattern, pattern))
        else:
            cur.execute("""
                SELECT p.id, p.name, p.internal_sku,
                       GROUP_CONCAT(e.sku, ', ')
                FROM products p
                LEFT JOIN external_skus e ON e.product_id = p.id
                GROUP BY p.id
                ORDER BY p.id DESC
            """)

        for row in cur.fetchall():
            product_id, name, internal, external = row
            self.tree.insert(
                "", "end",
                values=(product_id, name, internal, external or "—")
            )

    def get_selected_id(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showinfo("Выбор товара", "Сначала выберите товар.")
            return None
        return self.tree.item(selected[0], "values")[0]

    def show_product(self, event=None):
        product_id = self.get_selected_id()
        if product_id is None:
            return

        cur = self.conn.cursor()
        cur.execute(
            "SELECT name, internal_sku FROM products WHERE id = ?",
            (product_id,)
        )
        product = cur.fetchone()

        cur.execute(
            "SELECT sku FROM external_skus WHERE product_id = ? ORDER BY id",
            (product_id,)
        )
        external = [row[0] for row in cur.fetchall()]

        if not product:
            return

        dialog = tk.Toplevel(self.root)
        dialog.title("Товар")
        dialog.geometry("600x500")
        dialog.transient(self.root)
        dialog.grab_set()

        frame = ttk.Frame(dialog, padding=20)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(1, weight=1)
        frame.rowconfigure(2, weight=1)

        ttk.Label(frame, text="Название товара:").grid(
            row=0, column=0, sticky="w", pady=7, padx=(0, 10)
        )
        name_var = tk.StringVar(value=product[0])
        ttk.Entry(frame, textvariable=name_var).grid(
            row=0, column=1, sticky="ew", pady=7
        )

        ttk.Label(frame, text="Внутренний артикул:").grid(
            row=1, column=0, sticky="w", pady=7, padx=(0, 10)
        )
        internal_var = tk.StringVar(value=product[1])
        ttk.Entry(frame, textvariable=internal_var).grid(
            row=1, column=1, sticky="ew", pady=7
        )

        ttk.Label(frame, text="Внешние артикулы:").grid(
            row=2, column=0, sticky="nw", pady=7, padx=(0, 10)
        )

        ext_container = ttk.Frame(frame)
        ext_container.grid(row=2, column=1, sticky="nsew", pady=7)
        ext_container.columnconfigure(0, weight=1)
        ext_container.rowconfigure(0, weight=1)

        ext_list = ttk.Frame(ext_container)
        ext_list.grid(row=0, column=0, sticky="nsew")

        edit_entries = []

        def add_edit_entry(value=""):
            row = ttk.Frame(ext_list)
            row.pack(fill="x", pady=3)
            var = tk.StringVar(value=value)
            ttk.Entry(row, textvariable=var).pack(
                side="left", fill="x", expand=True
            )

            def remove():
                if len(edit_entries) <= 1:
                    var.set("")
                    return
                row.destroy()
                edit_entries[:] = [x for x in edit_entries if x[0] != row]

            ttk.Button(row, text="×", width=3, command=remove).pack(
                side="left", padx=(6, 0)
            )
            edit_entries.append((row, var))

        if external:
            for sku in external:
                add_edit_entry(sku)
        else:
            add_edit_entry()

        ttk.Button(
            ext_container,
            text="+ Добавить внешний артикул",
            command=add_edit_entry
        ).grid(row=1, column=0, sticky="w", pady=(8, 0))

        def update_product():
            name = name_var.get().strip()
            internal = internal_var.get().strip()

            external_values = []
            for _, var in edit_entries:
                value = var.get().strip()
                if value and value not in external_values:
                    external_values.append(value)

            if not name or not internal:
                messagebox.showwarning(
                    "Проверка",
                    "Название и внутренний артикул обязательны.",
                    parent=dialog
                )
                return

            try:
                cur = self.conn.cursor()
                cur.execute("""
                    UPDATE products
                    SET name = ?, internal_sku = ?
                    WHERE id = ?
                """, (name, internal, product_id))

                cur.execute(
                    "DELETE FROM external_skus WHERE product_id = ?",
                    (product_id,)
                )

                for sku in external_values:
                    cur.execute(
                        "INSERT INTO external_skus (product_id, sku) VALUES (?, ?)",
                        (product_id, sku)
                    )

                self.conn.commit()
                dialog.destroy()
                self.load_products()
                self.status_var.set(f"Товар ID {product_id} изменён.")

            except sqlite3.IntegrityError:
                self.conn.rollback()
                messagebox.showerror(
                    "Ошибка",
                    f"Внутренний артикул «{internal}» уже используется.",
                    parent=dialog
                )

        ttk.Button(
            frame, text="СОХРАНИТЬ ИЗМЕНЕНИЯ",
            style="Big.TButton",
            command=update_product
        ).grid(row=3, column=0, columnspan=2, sticky="w", pady=(20, 0))

    def delete_product(self):
        product_id = self.get_selected_id()
        if product_id is None:
            return

        selected = self.tree.selection()[0]
        values = self.tree.item(selected, "values")
        name = values[1]

        if not messagebox.askyesno(
            "Удаление",
            f"Удалить товар:\n\n{name}\n\nЭто также удалит все его внешние артикулы."
        ):
            return

        cur = self.conn.cursor()
        cur.execute("DELETE FROM external_skus WHERE product_id = ?", (product_id,))
        cur.execute("DELETE FROM products WHERE id = ?", (product_id,))
        self.conn.commit()

        self.load_products()
        self.status_var.set(f"Товар «{name}» удалён.")

    # -------------------- РАЗДЕЛ 2: EXCEL --------------------

    def build_excel_tab(self, parent):
        parent.columnconfigure(0, weight=1)
        parent.rowconfigure(2, weight=1)

        info = ttk.LabelFrame(parent, text=" 1. Выбор Excel-файлов ", padding=15)
        info.grid(row=0, column=0, sticky="ew", pady=(0, 10))
        info.columnconfigure(0, weight=1)

        ttk.Label(
            info,
            text="Можно загрузить любое количество Excel-файлов (.xlsx).\n"
                 "Из каждого файла используются столбцы «Товаров» и "
                 "«Артикул в системе партнера».\n"
                 "Один внешний артикул может быть указан у нескольких товаров: "
                 "для набора количество будет добавлено каждому товару из состава."
        ).grid(row=0, column=0, sticky="w", pady=(0, 10))

        files_frame = ttk.Frame(info)
        files_frame.grid(row=1, column=0, sticky="ew")
        files_frame.columnconfigure(0, weight=1)

        self.excel_listbox = tk.Listbox(
            files_frame, height=6, font=("Segoe UI", 10)
        )
        excel_files_scroll = ttk.Scrollbar(
            files_frame, orient="vertical", command=self.excel_listbox.yview
        )
        self.excel_listbox.configure(yscrollcommand=excel_files_scroll.set)
        self.excel_listbox.grid(row=0, column=0, sticky="nsew")
        excel_files_scroll.grid(row=0, column=1, sticky="ns")
        files_frame.rowconfigure(0, weight=1)

        files_buttons = ttk.Frame(files_frame)
        files_buttons.grid(row=0, column=2, sticky="ns", padx=(10, 0))

        ttk.Button(
            files_buttons,
            text="ДОБАВИТЬ EXCEL",
            command=self.add_excel_files
        ).pack(fill="x")

        ttk.Button(
            files_buttons,
            text="УБРАТЬ ВЫБРАННЫЙ",
            command=self.remove_excel_file
        ).pack(fill="x", pady=(7, 0))

        ttk.Button(
            files_buttons,
            text="ОЧИСТИТЬ",
            command=self.clear_excel_files
        ).pack(fill="x", pady=(7, 0))

        process = ttk.LabelFrame(parent, text=" 2. Обработка ", padding=15)
        process.grid(row=1, column=0, sticky="ew", pady=(0, 10))

        ttk.Button(
            process,
            text="СОЗДАТЬ ИТОГОВУЮ ТАБЛИЦУ",
            style="Big.TButton",
            command=self.process_excel
        ).pack(side="left")

        self.excel_status_var = tk.StringVar(
            value="Добавьте один или несколько Excel-файлов."
        )
        ttk.Label(
            process, textvariable=self.excel_status_var
        ).pack(side="left", padx=(15, 0))

        result_frame = ttk.LabelFrame(
            parent, text=" 3. Результат ", padding=10
        )
        result_frame.grid(row=2, column=0, sticky="nsew")
        result_frame.columnconfigure(0, weight=1)
        result_frame.rowconfigure(0, weight=1)

        columns = ("internal", "quantity")
        self.result_tree = ttk.Treeview(
            result_frame, columns=columns, show="headings"
        )
        self.result_tree.heading("internal", text="Внутренний артикул")
        self.result_tree.heading("quantity", text="Количество товаров")
        self.result_tree.column("internal", width=300)
        self.result_tree.column("quantity", width=220, anchor="center")

        result_scroll = ttk.Scrollbar(
            result_frame,
            orient="vertical",
            command=self.result_tree.yview
        )
        self.result_tree.configure(yscrollcommand=result_scroll.set)

        self.result_tree.grid(row=0, column=0, sticky="nsew")
        result_scroll.grid(row=0, column=1, sticky="ns")

        bottom = ttk.Frame(parent)
        bottom.grid(row=3, column=0, sticky="ew", pady=(10, 0))

        ttk.Button(
            bottom,
            text="СОХРАНИТЬ РЕЗУЛЬТАТ В EXCEL",
            style="Big.TButton",
            command=self.save_excel_result
        ).pack(side="left")

    def add_excel_files(self):
        if load_workbook is None:
            messagebox.showerror(
                "Не установлен openpyxl",
                "Для работы с Excel установите библиотеку:\n\n"
                "python -m pip install openpyxl"
            )
            return

        files = filedialog.askopenfilenames(
            title="Выберите Excel-файлы",
            filetypes=[
                ("Excel files", "*.xlsx"),
                ("Все файлы", "*.*")
            ]
        )

        for path in files:
            if path not in self.excel_files:
                self.excel_files.append(path)

        self.refresh_excel_file_list()

    def remove_excel_file(self):
        selection = self.excel_listbox.curselection()
        if not selection:
            return

        index = selection[0]
        del self.excel_files[index]
        self.refresh_excel_file_list()

    def clear_excel_files(self):
        self.excel_files = []
        self.excel_result = []

        for item in self.result_tree.get_children():
            self.result_tree.delete(item)

        self.refresh_excel_file_list()

    def refresh_excel_file_list(self):
        self.excel_listbox.delete(0, tk.END)

        for path in self.excel_files:
            self.excel_listbox.insert(tk.END, os.path.basename(path))

        if not self.excel_files:
            self.excel_status_var.set("Добавьте один или несколько Excel-файлов.")
        else:
            self.excel_status_var.set(
                f"Выбрано файлов: {len(self.excel_files)}"
            )

    @staticmethod
    def normalize_header(value):
        if value is None:
            return ""
        return " ".join(str(value).strip().lower().split())

    @staticmethod
    def normalize_sku(value):
        if value is None:
            return ""

        # Excel иногда возвращает числовые артикулы как 12345.0
        if isinstance(value, float) and value.is_integer():
            return str(int(value))

        text = str(value).strip()
        if text.endswith(".0") and text[:-2].isdigit():
            text = text[:-2]

        return text

    @staticmethod
    def parse_quantity(value):
        if value is None or str(value).strip() == "":
            return 0

        if isinstance(value, bool):
            return int(value)

        if isinstance(value, (int, float)):
            return float(value)

        text = str(value).strip().replace(" ", "").replace(",", ".")

        try:
            return float(text)
        except ValueError:
            return None

    def find_column(self, headers, wanted):
        wanted_norm = self.normalize_header(wanted)

        # Сначала точное совпадение
        for i, header in enumerate(headers):
            if self.normalize_header(header) == wanted_norm:
                return i

        # Затем мягкое совпадение: чтобы пережить лишние пробелы/символы
        for i, header in enumerate(headers):
            h = self.normalize_header(header)
            if wanted_norm in h or h in wanted_norm:
                return i

        return None

    def read_excel_file(self, path):
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active

        rows = sheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            workbook.close()
            raise ValueError("Файл пустой.")

        headers = list(headers)

        quantity_col = self.find_column(headers, "Товаров")
        partner_col = self.find_column(
            headers, "Артикул в системе партнера"
        )

        if quantity_col is None or partner_col is None:
            workbook.close()

            found = ", ".join(
                str(x).strip() for x in headers
                if x is not None and str(x).strip()
            )

            missing = []
            if quantity_col is None:
                missing.append("«Товаров»")
            if partner_col is None:
                missing.append("«Артикул в системе партнера»")

            raise ValueError(
                "Не найдены столбцы: " + ", ".join(missing) +
                f"\n\nСтолбцы, которые найдены в файле:\n{found}"
            )

        data = []

        for row_number, row in enumerate(rows, start=2):
            row = list(row)

            max_index = max(quantity_col, partner_col)
            if len(row) <= max_index:
                continue

            partner_sku = self.normalize_sku(row[partner_col])
            quantity = self.parse_quantity(row[quantity_col])

            if not partner_sku:
                continue

            if quantity is None:
                raise ValueError(
                    f"Строка {row_number}: значение в «Товаров» "
                    f"не является числом: {row[quantity_col]!r}"
                )

            if quantity == 0:
                continue

            data.append((partner_sku, quantity))

        workbook.close()
        return data

    def get_sku_mapping(self):
        """
        Возвращает соответствия внешний артикул -> список внутренних артикулов.

        Один внешний артикул может принадлежать нескольким товарам. Это нужно
        для наборов: если внешний артикул набора встречается в Excel с
        количеством 5, и он указан у трёх товаров, то каждому из этих
        внутренних артикулов будет начислено по 5 штук.
        """
        cur = self.conn.cursor()
        cur.execute("""
            SELECT p.internal_sku, e.sku
            FROM products p
            JOIN external_skus e ON e.product_id = p.id
            ORDER BY p.id, e.id
        """)

        mapping = {}

        for internal, external in cur.fetchall():
            internal = self.normalize_sku(internal)
            external = self.normalize_sku(external)

            if not external or not internal:
                continue

            mapping.setdefault(external, [])
            if internal not in mapping[external]:
                mapping[external].append(internal)

        return mapping

    def process_excel(self):
        if load_workbook is None:
            messagebox.showerror(
                "Не установлен openpyxl",
                "Установите библиотеку командой:\n"
                "python -m pip install openpyxl"
            )
            return

        if not self.excel_files:
            messagebox.showwarning(
                "Нет файлов",
                "Сначала добавьте хотя бы один Excel-файл."
            )
            return

        mapping = self.get_sku_mapping()

        if not mapping:
            messagebox.showwarning(
                "Нет товаров",
                "В базе пока нет внешних артикулов.\n"
                "Сначала создайте товары в разделе «Товары»."
            )
            return

        totals = {}
        not_found = set()
        errors = []

        for path in self.excel_files:
            try:
                rows = self.read_excel_file(path)
            except Exception as exc:
                errors.append(
                    f"{os.path.basename(path)}:\n{exc}"
                )
                continue

            for partner_sku, quantity in rows:
                internals = mapping.get(partner_sku)

                if internals is None:
                    not_found.add(partner_sku)
                    continue

                # Один внешний артикул может быть артикулом набора,
                # состоящего из нескольких товаров. В этом случае количество
                # набора добавляется КАЖДОМУ товару из состава набора.
                for internal in internals:
                    totals[internal] = totals.get(internal, 0) + quantity

        if errors:
            messagebox.showerror(
                "Ошибка обработки",
                "\n\n".join(errors[:10]) +
                ("\n\n... и другие ошибки." if len(errors) > 10 else "")
            )
            return

        # Убираем .0 для целых количеств
        self.excel_result = []
        for internal in sorted(totals.keys()):
            quantity = totals[internal]
            if float(quantity).is_integer():
                quantity = int(quantity)
            self.excel_result.append((internal, quantity))

        for item in self.result_tree.get_children():
            self.result_tree.delete(item)

        for internal, quantity in self.excel_result:
            self.result_tree.insert(
                "", "end", values=(internal, quantity)
            )

        status = (
            f"Готово. Найдено внутренних артикулов: "
            f"{len(self.excel_result)}"
        )

        if not_found:
            status += f" | Не найдено внешних артикулов: {len(not_found)}"

        self.excel_status_var.set(status)

        if not_found:
            not_found_text = "\n".join(
                sorted(not_found, key=str)[:30]
            )
            more = (
                f"\n... и ещё {len(not_found) - 30}"
                if len(not_found) > 30 else ""
            )

            messagebox.showwarning(
                "Обработка завершена",
                f"Готово.\n\n"
                f"В итоговую таблицу попало товаров: "
                f"{len(self.excel_result)}.\n\n"
                f"Не найдены в базе внешние артикулы: "
                f"{len(not_found)}.\n\n"
                f"{not_found_text}{more}"
            )
        else:
            messagebox.showinfo(
                "Готово",
                f"Обработка завершена.\n\n"
                f"В итоговой таблице: {len(self.excel_result)} "
                f"внутренних артикулов."
            )

    @staticmethod
    def get_downloads_folder():
        """Возвращает папку «Загрузки» текущего пользователя."""
        home = os.path.expanduser("~")
        candidates = []

        if os.name == "nt":
            userprofile = os.environ.get("USERPROFILE", home)
            candidates.append(os.path.join(userprofile, "Downloads"))
            candidates.append(os.path.join(userprofile, "Загрузки"))
        else:
            xdg_download = os.environ.get("XDG_DOWNLOAD_DIR")
            if xdg_download:
                candidates.append(os.path.expanduser(xdg_download))
            candidates.append(os.path.join(home, "Downloads"))
            candidates.append(os.path.join(home, "Загрузки"))

        for path in candidates:
            if path and os.path.isdir(path):
                return path

        downloads = os.path.join(home, "Downloads")
        os.makedirs(downloads, exist_ok=True)
        return downloads

    @staticmethod
    def unique_download_path(folder, filename="Итог.xlsx"):
        """
        Итог.xlsx, если свободен.
        Иначе Итог(1).xlsx, Итог(2).xlsx и т.д.
        """
        name, ext = os.path.splitext(filename)
        path = os.path.join(folder, filename)
        if not os.path.exists(path):
            return path

        index = 1
        while True:
            candidate = os.path.join(folder, f"{name}({index}){ext}")
            if not os.path.exists(candidate):
                return candidate
            index += 1

    def save_excel_result(self):
        if load_workbook is None or Workbook is None:
            messagebox.showerror(
                "Не установлен openpyxl",
                "Установите библиотеку:\npython -m pip install openpyxl"
            )
            return

        if not self.excel_result:
            messagebox.showwarning(
                "Нет результата",
                "Сначала создайте итоговую таблицу."
            )
            return

        downloads = self.get_downloads_folder()
        path = self.unique_download_path(downloads, "Итог.xlsx")

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Результат"

        sheet["A1"] = "Внутренний артикул"
        sheet["B1"] = "Количество товаров"

        for row_index, (internal, quantity) in enumerate(
            self.excel_result, start=2
        ):
            sheet.cell(row=row_index, column=1, value=internal)
            sheet.cell(row=row_index, column=2, value=quantity)

        sheet.column_dimensions["A"].width = 25
        sheet.column_dimensions["B"].width = 22

        try:
            workbook.save(path)
        except OSError as exc:
            messagebox.showerror(
                "Ошибка",
                f"Не удалось сохранить файл:\n\n{exc}"
            )
            return

        self.excel_status_var.set(f"Сохранено: {os.path.basename(path)}")
        messagebox.showinfo(
            "Сохранено",
            f"Итоговая таблица сохранена в «Загрузки»:\n\n{path}"
        )

    # -------------------- РАЗДЕЛ 3: PDF --------------------

    def build_pdf_tab(self, parent):
        bg = "#f4f5f7"
        shell = tk.Frame(parent, bg=bg)
        shell.pack(fill="both", expand=True)

        tk.Label(
            shell,
            text="Сортировщик чеков Kaspi",
            font=("Segoe UI", 24, "bold"),
            bg=bg,
            fg="#202124",
        ).pack(pady=(22, 4))

        tk.Label(
            shell,
            text="Перетащите папку с чеками сюда, нажмите «Выбрать папку» или Ctrl + V",
            font=("Segoe UI", 11),
            bg=bg,
            fg="#6b7280",
        ).pack(pady=(0, 16))

        drop = tk.Frame(
            shell,
            bg="#ffffff",
            highlightbackground="#b9bec8",
            highlightthickness=2,
            cursor="hand2",
        )
        drop.pack(padx=70, fill="x")

        drop_label = tk.Label(
            drop,
            text="📁\n\nПЕРЕТАЩИТЕ ПАПКУ СЮДА",
            font=("Segoe UI", 16, "bold"),
            bg="#ffffff",
            fg="#4b5563",
            justify="center",
            cursor="hand2",
        )
        drop_label.pack(pady=(28, 8))

        drop_hint = tk.Label(
            drop,
            text="или скопируйте папку в Проводнике → Ctrl + V",
            font=("Segoe UI", 10),
            bg="#ffffff",
            fg="#9ca3af",
            cursor="hand2",
        )
        drop_hint.pack(pady=(0, 12))

        choose = tk.Button(
            drop,
            text="ВЫБРАТЬ ПАПКУ",
            command=self.pdf_choose_folder,
            font=("Segoe UI", 11, "bold"),
            bg="#111827",
            fg="white",
            activebackground="#374151",
            activeforeground="white",
            relief="flat",
            cursor="hand2",
            padx=16,
            pady=8,
        )
        choose.pack(pady=(0, 20))

        for widget in (drop, drop_label, drop_hint):
            widget.bind("<Button-1>", lambda _e: self.pdf_choose_folder())

        self.pdf_enable_drop(drop, drop_label, drop_hint, choose)

        info = tk.Frame(shell, bg=bg)
        info.pack(fill="x", padx=80, pady=18)

        self.pdf_folder_var = tk.StringVar(value="Папка не выбрана")
        self.pdf_files_var = tk.StringVar(value="")
        self.pdf_products_var = tk.StringVar(value="")
        self.pdf_status_var = tk.StringVar(value="Ожидание папки...")

        tk.Label(
            info,
            textvariable=self.pdf_folder_var,
            font=("Segoe UI", 12, "bold"),
            bg=bg,
            fg="#202124",
            anchor="w",
        ).pack(fill="x")
        tk.Label(
            info,
            textvariable=self.pdf_files_var,
            font=("Segoe UI", 11),
            bg=bg,
            fg="#4b5563",
            anchor="w",
        ).pack(fill="x", pady=(5, 0))
        tk.Label(
            info,
            textvariable=self.pdf_products_var,
            font=("Segoe UI", 11),
            bg=bg,
            fg="#4b5563",
            anchor="w",
        ).pack(fill="x", pady=(3, 0))

        self.pdf_status_label = tk.Label(
            shell,
            textvariable=self.pdf_status_var,
            font=("Segoe UI", 10),
            bg=bg,
            fg="#6b7280",
        )
        self.pdf_status_label.pack(pady=(0, 8))

        buttons = tk.Frame(shell, bg=bg)
        buttons.pack(pady=(0, 20))

        self.pdf_action_button = tk.Button(
            buttons,
            text="СОЗДАТЬ",
            command=self.pdf_on_action,
            font=("Segoe UI", 12, "bold"),
            width=18,
            height=2,
            bg="#111827",
            fg="white",
            activebackground="#374151",
            activeforeground="white",
            relief="flat",
            cursor="hand2",
            state="disabled",
        )
        self.pdf_action_button.pack(side="left", padx=8)

        self.pdf_reset_button = tk.Button(
            buttons,
            text="СБРОС",
            command=self.pdf_reset,
            font=("Segoe UI", 12, "bold"),
            width=18,
            height=2,
            bg="#e5e7eb",
            fg="#111827",
            activebackground="#d1d5db",
            relief="flat",
            cursor="hand2",
        )
        self.pdf_reset_button.pack(side="left", padx=8)

        for widget in (shell, drop, drop_label, drop_hint, info, buttons):
            widget.bind("<Control-v>", self.pdf_on_ctrl_v)
            widget.bind("<Control-V>", self.pdf_on_ctrl_v)
        parent.bind("<Control-v>", self.pdf_on_ctrl_v)
        parent.bind("<Control-V>", self.pdf_on_ctrl_v)

    def pdf_enable_drop(self, *widgets):
        if DND_FILES is None:
            return
        if not hasattr(self.root, "drop_target_register"):
            return
        for widget in widgets:
            try:
                widget.drop_target_register(DND_FILES)
                widget.dnd_bind("<<Drop>>", self.pdf_on_drop)
            except tk.TclError:
                pass

    def pdf_on_drop(self, event):
        if self.pdf_processing:
            return
        raw = getattr(event, "data", "") or ""
        try:
            paths = self.root.tk.splitlist(raw)
        except tk.TclError:
            paths = [raw.strip("{}")]
        if not paths:
            return
        folder = _folder_if_valid(paths[0])
        if folder:
            self.pdf_select_folder(folder)
        else:
            messagebox.showwarning(
                "Нужна папка",
                "Перетащите именно папку с PDF-чеками.",
            )

    def pdf_choose_folder(self):
        if self.pdf_processing:
            return
        folder = filedialog.askdirectory(
            parent=self.root,
            title="Папка с PDF-чеками",
            mustexist=True,
        )
        if folder:
            self.pdf_select_folder(Path(folder))

    def pdf_paste_folder(self):
        self.pdf_on_ctrl_v()

    def pdf_tab_selected(self):
        try:
            return str(self.notebook.select()) == str(self.pdf_tab)
        except Exception:
            return False

    def pdf_on_ctrl_v(self, event=None):
        if not self.pdf_tab_selected():
            return
        widget = getattr(event, "widget", None) if event else None
        try:
            class_name = widget.winfo_class() if widget else ""
        except tk.TclError:
            class_name = ""
        if class_name in ("TEntry", "Entry", "Text", "TCombobox"):
            return
        if self.pdf_processing:
            return "break"
        folder = get_folder_from_clipboard(self.root)
        if folder:
            self.pdf_select_folder(folder)
        else:
            messagebox.showinfo(
                "Папка не найдена",
                "Скопируйте папку в Проводнике через Ctrl+C,\n"
                "затем на вкладке «Сортировка PDF» нажмите Ctrl+V.\n"
                "Либо нажмите «Выбрать папку».",
            )
        return "break"

    def pdf_select_folder(self, folder):
        if fitz is None:
            messagebox.showerror(
                "Нет PyMuPDF",
                "Для сортировки PDF нужна библиотека pymupdf.\n"
                "Установите новую сборку через ProductManagerSetup.exe.",
            )
            return

        folder = Path(folder)
        pdf_files = [
            path for path in folder.iterdir()
            if path.is_file() and path.suffix.lower() == ".pdf"
        ]
        if not pdf_files:
            messagebox.showwarning(
                "PDF не найдены",
                "В выбранной папке нет PDF-файлов.",
            )
            return

        self.pdf_selected_folder = folder
        self.pdf_files = sorted(pdf_files)
        self.pdf_result_file = None
        self.pdf_folder_var.set(f"Папка: {folder}")
        self.pdf_files_var.set(f"PDF-файлов: {len(pdf_files)}")
        self.pdf_products_var.set("Считывание чеков...")
        self.pdf_status_var.set("Считываю содержимое чеков...")
        self.pdf_action_button.config(state="disabled", text="СЧИТЫВАНИЕ...")
        self.pdf_processing = True

        threading.Thread(target=self.pdf_scan_folder, daemon=True).start()

    def pdf_scan_folder(self):
        single = []
        multiple = []
        unknown = []
        for pdf_path in self.pdf_files:
            products = extract_receipt_products(pdf_path)
            if len(products) == 1:
                name, quantity = products[0]
                single.append((name, quantity, pdf_path))
            elif len(products) > 1:
                multiple.append((products, pdf_path))
            else:
                unknown.append(pdf_path)
        single.sort(key=lambda item: (item[0].lower(), -item[1]))
        self.root.after(
            0, lambda: self.pdf_scan_finished(single, multiple, unknown)
        )

    def pdf_scan_finished(self, single, multiple, unknown):
        self.pdf_single = single
        self.pdf_multiple = multiple
        self.pdf_unknown = unknown
        self.pdf_processing = False
        unique_products = len({name for name, _qty, _path in single})
        self.pdf_products_var.set(
            f"Товаров найдено: {unique_products}  |  "
            f"Один товар: {len(single)} чеков  |  "
            f"Несколько товаров: {len(multiple)}  |  "
            f"Нераспознанные: {len(unknown)}"
        )
        self.pdf_status_var.set("Проверка завершена. Можно создавать итоговый PDF.")
        self.pdf_action_button.config(state="normal", text="СОЗДАТЬ")

    def pdf_on_action(self):
        if self.pdf_result_file:
            self.pdf_download_result()
        else:
            self.pdf_create_result()

    def pdf_create_result(self):
        if not self.pdf_selected_folder or self.pdf_processing:
            return
        self.pdf_processing = True
        self.pdf_action_button.config(state="disabled", text="СОЗДАНИЕ...")
        self.pdf_status_var.set("Создаю итоговый PDF...")
        threading.Thread(target=self.pdf_create_thread, daemon=True).start()

    def pdf_create_thread(self):
        folder = self.pdf_selected_folder
        output_file = folder.parent / (folder.name + ".pdf")
        if output_file.exists():
            try:
                output_file.unlink()
            except OSError:
                self.root.after(
                    0,
                    lambda: messagebox.showerror(
                        "Ошибка",
                        "Не удалось заменить существующий итоговый PDF.\n"
                        "Закройте его, если он открыт.",
                    ),
                )
                self.root.after(0, self.pdf_reset_after_error)
                return

        output = fitz.open()
        try:
            for _name, _qty, pdf_path in self.pdf_single:
                source = fitz.open(pdf_path)
                output.insert_pdf(source)
                source.close()
            for _products, pdf_path in self.pdf_multiple:
                source = fitz.open(pdf_path)
                output.insert_pdf(source)
                source.close()
            for pdf_path in self.pdf_unknown:
                source = fitz.open(pdf_path)
                output.insert_pdf(source)
                source.close()
            output.save(output_file, garbage=4, deflate=True)
            output.close()
            self.root.after(0, lambda: self.pdf_creation_finished(output_file))
        except Exception as exc:
            try:
                output.close()
            except Exception:
                pass
            self.root.after(
                0,
                lambda e=exc: messagebox.showerror("Ошибка создания PDF", str(e)),
            )
            self.root.after(0, self.pdf_reset_after_error)

    def pdf_creation_finished(self, output_file):
        self.pdf_processing = False
        self.pdf_result_file = output_file
        self.pdf_status_var.set("Готово. Итоговый PDF создан.")
        self.pdf_action_button.config(state="normal", text="СКАЧАТЬ")
        messagebox.showinfo(
            "Готово",
            f"Файл создан:\n\n{output_file}\n\n"
            "Нажмите «СКАЧАТЬ», чтобы скопировать его в «Загрузки».",
        )

    def pdf_reset_after_error(self):
        self.pdf_processing = False
        self.pdf_action_button.config(
            state="normal" if self.pdf_selected_folder else "disabled",
            text="СОЗДАТЬ",
        )
        self.pdf_status_var.set("Готово к новой обработке.")

    def pdf_download_result(self):
        if not self.pdf_result_file or not Path(self.pdf_result_file).exists():
            messagebox.showerror("Ошибка", "Итоговый PDF не найден.")
            return
        downloads = Path(self.get_downloads_folder())
        downloads.mkdir(parents=True, exist_ok=True)
        destination = Path(
            self.unique_download_path(str(downloads), Path(self.pdf_result_file).name)
        )
        try:
            shutil.copy2(self.pdf_result_file, destination)
        except OSError as exc:
            messagebox.showerror("Ошибка", str(exc))
            return
        self.pdf_status_var.set(f"Сохранено в «Загрузки»: {destination.name}")
        messagebox.showinfo("Готово", f"Файл скопирован в:\n\n{destination}")

    def pdf_reset(self):
        if self.pdf_processing:
            return
        self.pdf_selected_folder = None
        self.pdf_result_file = None
        self.pdf_files = []
        self.pdf_single = []
        self.pdf_multiple = []
        self.pdf_unknown = []
        self.pdf_folder_var.set("Папка не выбрана")
        self.pdf_files_var.set("")
        self.pdf_products_var.set("")
        self.pdf_status_var.set("Ожидание папки с PDF-чеками.")
        self.pdf_action_button.config(state="disabled", text="СОЗДАТЬ")

    # -------------------- ОБНОВЛЕНИЯ --------------------

    def check_for_updates_async(self):
        if self._update_busy:
            return
        repo = get_github_repo()
        if not repo or "YOUR_GITHUB" in repo:
            self.update_status_var.set("Репозиторий не указан")
            return
        self.update_status_var.set("Проверка обновлений...")
        self.update_button.config(state="disabled")
        threading.Thread(target=self._check_for_updates, daemon=True).start()

    def _github_json(self, url):
        req = urllib.request.Request(
            url,
            headers={
                "User-Agent": f"ProductManager/{APP_VERSION}",
                "Accept": "application/vnd.github+json",
            }
        )
        with urllib.request.urlopen(req, timeout=20) as resp:
            return json.loads(resp.read().decode("utf-8"))

    def _read_text_url(self, url):
        req = urllib.request.Request(
            url,
            headers={"User-Agent": f"ProductManager/{APP_VERSION}"}
        )
        with urllib.request.urlopen(req, timeout=20) as resp:
            return resp.read().decode("utf-8").strip()

    def _check_branch_version(self, repo):
        version_url = (
            f"https://raw.githubusercontent.com/{repo}/"
            f"{GITHUB_BRANCH}/version.txt"
        )
        script_url = (
            f"https://raw.githubusercontent.com/{repo}/"
            f"{GITHUB_BRANCH}/{UPDATE_SCRIPT_NAME}"
        )
        latest = self._read_text_url(version_url).splitlines()[0].strip()
        return {
            "version": latest,
            "url": script_url,
            "notes": "Обновление кода с GitHub (.py).",
        }

    def _check_for_updates(self):
        repo = get_github_repo()
        try:
            info = self._check_branch_version(repo)
            self.root.after(0, lambda: self._apply_update_info(info))
        except urllib.error.HTTPError as exc:
            if exc.code == 404:
                msg = "На GitHub нет version.txt или .py"
            else:
                msg = f"GitHub ответил ошибкой {exc.code}"
            self.root.after(0, lambda m=msg: self._set_update_idle(m))
        except (urllib.error.URLError, TimeoutError, OSError):
            self.root.after(0, lambda: self._set_update_idle("Нет связи с GitHub"))

    def _set_update_idle(self, text):
        self.latest_update = None
        self.update_status_var.set(text)
        self.update_status_label.configure(style="Version.TLabel")
        self.update_button.config(state="disabled")

    def _apply_update_info(self, info):
        latest = info.get("version") or ""
        if not latest:
            self._set_update_idle("Не удалось прочитать версию")
            return

        if not version_newer(latest, APP_VERSION):
            self._set_update_idle("Установлена последняя версия")
            return

        self.latest_update = info
        self.update_status_var.set(f"Доступна {latest}")
        self.update_status_label.configure(style="Update.TLabel")
        if info.get("url") and getattr(sys, "frozen", False):
            self.update_button.config(state="normal")
        elif info.get("url"):
            self.update_button.config(state="normal")
        else:
            self.update_button.config(state="disabled")

    def start_update(self):
        info = self.latest_update
        if not info or not info.get("url"):
            messagebox.showinfo("Обновление", "Сначала проверьте наличие обновления.")
            return

        if self._update_busy:
            return

        notes = info.get("notes") or "Новая версия программы."
        if len(notes) > 600:
            notes = notes[:600] + "..."
        if not messagebox.askyesno(
            "Обновление",
            f"Установить версию {info.get('version')}?\n\n"
            f"{notes}\n\n"
            "Скачается код (.py), рабочее окружение сохранится в AppData.\n"
            "База products.db тоже будет в скрытой папке."
        ):
            return

        self._update_busy = True
        self.update_button.config(state="disabled")
        self.update_status_var.set("Скачивание...")
        threading.Thread(
            target=self._download_and_apply_update,
            args=(info,),
            daemon=True
        ).start()

    def _download_and_apply_update(self, info):
        dest = os.path.join(data_dir(), LOCAL_UPDATE_SCRIPT)
        temp_dest = dest + ".tmp"
        version_path = os.path.join(data_dir(), "version.txt")
        try:
            req = urllib.request.Request(
                info["url"],
                headers={"User-Agent": f"ProductManager/{APP_VERSION}"}
            )
            with urllib.request.urlopen(req, timeout=60) as resp:
                data = resp.read()
            text = data.decode("utf-8")
            if "class ProductApp" not in text or len(text) < 1000:
                raise OSError("Скачанный файл не похож на программу")
            with open(temp_dest, "w", encoding="utf-8", newline="\n") as f:
                f.write(text)
            os.replace(temp_dest, dest)
            with open(version_path, "w", encoding="utf-8") as f:
                f.write(str(info.get("version") or "").strip() + "\n")
            persist_runtime()
        except Exception as exc:
            try:
                if os.path.exists(temp_dest):
                    os.remove(temp_dest)
            except OSError:
                pass
            self.root.after(0, lambda e=exc: self._update_failed(e))
            return

        self.root.after(0, self._restart_after_script_update)

    def _update_failed(self, exc):
        self._update_busy = False
        self.update_status_var.set("Ошибка обновления")
        if self.latest_update:
            self.update_button.config(state="normal")
        messagebox.showerror("Обновление", f"Не удалось скачать обновление:\n\n{exc}")

    def _restart_after_script_update(self):
        self.update_status_var.set("Перезапуск...")
        try:
            if getattr(sys, "frozen", False):
                target = f'"{sys.executable}"'
            else:
                target = f'"{sys.executable}" "{os.path.abspath(__file__)}"'

            bat_path = os.path.join(tempfile.gettempdir(), "pm_restart.bat")
            bat = (
                "@echo off\r\n"
                "ping 127.0.0.1 -n 3 >nul\r\n"
                f'start "" {target}\r\n'
                "del \"%~f0\"\r\n"
            )
            with open(bat_path, "w", encoding="ascii", newline="\r\n") as f:
                f.write(bat)

            kwargs = {"cwd": tempfile.gettempdir(), "close_fds": True}
            if os.name == "nt":
                kwargs["creationflags"] = (
                    getattr(subprocess, "DETACHED_PROCESS", 0x00000008)
                    | getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0x00000200)
                )
            subprocess.Popen(["cmd.exe", "/c", bat_path], **kwargs)
        except OSError as exc:
            self._update_failed(exc)
            return
        self.close()

    def close(self):
        self.conn.close()
        self.root.destroy()


def _run_downloaded_script():
    """Exe запускает свежий .py из папки, если он уже скачан с GitHub."""
    if os.environ.get("PM_BOOTSTRAPPED") == "1":
        return False
    if not getattr(sys, "frozen", False):
        return False
    updated = os.path.join(data_dir(), LOCAL_UPDATE_SCRIPT)
    if not os.path.isfile(updated):
        legacy = os.path.join(app_dir(), LOCAL_UPDATE_SCRIPT)
        if os.path.isfile(legacy):
            try:
                os.replace(legacy, updated)
            except OSError:
                updated = legacy
        else:
            return False
    if not version_newer(update_app_version(), bundled_app_version()):
        return False
    os.environ["PM_BOOTSTRAPPED"] = "1"
    import runpy
    try:
        runpy.run_path(updated, run_name="__main__")
        return True
    except Exception:
        try:
            os.remove(updated)
        except OSError:
            pass
        return False


if __name__ == "__main__":
    if not _run_downloaded_script():
        cleanup_app_folder()
        root = TkinterDnD.Tk() if TkinterDnD is not None else tk.Tk()
        app = ProductApp(root)
        root.mainloop()
